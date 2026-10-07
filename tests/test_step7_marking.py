"""Tests of step 7, the marking file (decision 0035). No test makes a model call.

Run from the folder ccar-f-mock-generator:
    python3 -m unittest discover tests
"""
import json
import os
import shutil
import sys
import tempfile
import unittest

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))

from step1_plan import make_plan, RECORDS as RECORDS_PATH  # noqa: E402
from step7_marking import Marking, MarkingError, HEADERS, plain  # noqa: E402

with open(RECORDS_PATH, encoding="utf-8") as f:
    RECORDS = json.load(f)
PLAN = make_plan(RECORDS, 1, 11, "2026-10-07")


def question_for(slot):
    options = [{"letter": "ABCDE"[i], "text": "Option %d of %s." % (i, slot["slot_id"]),
                "correct": i < slot["select"], "reason": "Reason %d." % i} for i in range(slot["options"])]
    return {"question_id": slot["slot_id"] + "-T1", "slot_id": slot["slot_id"], "try": 1, "point": slot["point"],
            "select": slot["select"], "status": "written", "question": "Question of %s?" % slot["slot_id"],
            "options": options}


class TestMarking(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.m = Marking(PLAN, RECORDS, self.folder)
        self.questions = [question_for(s) for s in PLAN["slots"]]

    def tearDown(self):
        shutil.rmtree(self.folder)

    def fill(self, round_number, values):
        """Act as the owner: write marks into the file."""
        path = self.m.path(round_number)
        wb = load_workbook(path)
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            qid = row[0].value
            if qid in values:
                mark, reason = values[qid]
                row[HEADERS.index("Mark")].value = mark
                row[HEADERS.index("Reason (needed for a bad mark)")].value = reason
        wb.save(path)

    def test_one_row_for_each_question_and_the_columns(self):
        path = self.m.make(1, self.questions)
        self.assertEqual(os.path.basename(path), "M1-round1.xlsx")
        ws = load_workbook(path).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), HEADERS)
        self.assertEqual(len(rows), 61)
        first = dict(zip(HEADERS, rows[1]))
        slot = PLAN["slots"][0]
        self.assertEqual(first["Question ID"], slot["slot_id"] + "-T1")
        self.assertEqual(first["Select"], "Select %d." % slot["select"])
        self.assertEqual(first["Point"], slot["point"])
        point = [p for t in RECORDS["task_statements"] for p in t["points"] if p["id"] == slot["point"]][0]
        self.assertEqual(first["Point text"], plain(point["text"]))
        self.assertIn("A: Reason 0.", first["Reasons of the options"])

    def test_the_file_shows_no_verdict(self):
        path = self.m.make(1, self.questions)
        rows = [list(r) for r in load_workbook(path).active.iter_rows(values_only=True)]
        for header in rows[0]:
            for word in ("verdict", "finding", "step 3", "step 4", "step 6"):
                self.assertNotIn(word, header.lower())
        cells = {str(c).strip().lower() for r in rows for c in r if c is not None}
        self.assertFalse(cells & {"pass", "fail", "cannot_decide", "cannot decide"})

    def test_the_two_columns_of_the_owner_are_yellow_and_mark_has_a_list(self):
        ws = load_workbook(self.m.make(1, self.questions)).active
        mark_col = HEADERS.index("Mark")
        self.assertEqual(ws.cell(row=2, column=mark_col + 1).fill.fgColor.rgb[-6:], "FFF2CC")
        self.assertEqual(ws.cell(row=2, column=mark_col + 2).fill.fgColor.rgb[-6:], "FFF2CC")
        self.assertEqual(ws.cell(row=2, column=1).fill.fgColor.rgb[-6:], "000000")
        formulas = [v.formula1 for v in ws.data_validations.dataValidation]
        self.assertIn('"good,bad"', formulas)

    def test_read_the_marks(self):
        self.m.make(1, self.questions[:3])
        ids = [q["question_id"] for q in self.questions[:3]]
        self.fill(1, {ids[0]: ("good", None), ids[1]: (" Bad ", "C is also correct.")})
        marks = self.m.read(1)
        self.assertEqual(marks, [{"question_id": ids[0], "mark": "good", "reason": ""},
                                 {"question_id": ids[1], "mark": "bad", "reason": "C is also correct."},
                                 {"question_id": ids[2], "mark": None, "reason": ""}])

    def test_no_file_with_marks_is_written_over(self):
        self.m.make(1, self.questions[:2])
        self.fill(1, {self.questions[0]["question_id"]: ("good", None)})
        with self.assertRaises(MarkingError):
            self.m.make(1, self.questions[:2])

    def test_the_same_empty_file_after_a_stop_is_used_again(self):
        first = self.m.make(1, self.questions[:2])
        self.assertEqual(self.m.make(1, self.questions[:2]), first)
        with self.assertRaises(MarkingError):
            self.m.make(1, self.questions[:3])

    def test_a_question_two_times_is_an_error(self):
        path = self.m.make(1, self.questions[:2])
        wb = load_workbook(path)
        ws = wb.active
        ws.append([c.value for c in ws[2]])
        wb.save(path)
        with self.assertRaises(MarkingError) as e:
            self.m.read(1)
        self.assertIn("two times", str(e.exception))

    def test_a_formula_is_an_error(self):
        self.m.make(1, self.questions[:2])
        self.fill(1, {self.questions[0]["question_id"]: ('=IF(1,"bad","good")', None)})
        with self.assertRaises(MarkingError) as e:
            self.m.read(1)
        self.assertIn("formula", str(e.exception))

    def test_another_sheet_first_still_reads_the_marks(self):
        path = self.m.make(1, self.questions[:2])
        wb = load_workbook(path)
        wb.create_sheet("Notes", 0)
        wb.active = 0
        wb.save(path)
        self.assertEqual(len(self.m.read(1)), 2)

    def test_the_sheet_is_protected_and_only_the_owner_columns_are_open(self):
        ws = load_workbook(self.m.make(1, self.questions[:2])).active
        self.assertTrue(ws.protection.sheet)
        self.assertTrue(ws.protection.sort)   # True means: sorting is not permitted
        mark_col = HEADERS.index("Mark") + 1
        self.assertFalse(ws.cell(row=2, column=mark_col).protection.locked)
        self.assertFalse(ws.cell(row=2, column=mark_col + 1).protection.locked)
        self.assertTrue(ws.cell(row=2, column=1).protection.locked)
        v = ws.data_validations.dataValidation[0]
        self.assertTrue(v.showErrorMessage)
        self.assertEqual(str(v.sqref), "N2:N3")

    def test_a_text_that_starts_with_an_equals_sign_is_text(self):
        q = self.questions[0]
        q["options"][0]["text"] = "= is the operator that assigns a value."
        q["options"][1]["text"] = "A text with a bell \x07 character."
        ws = load_workbook(self.m.make(1, [q])).active
        self.assertEqual(ws.cell(row=2, column=HEADERS.index("A") + 1).data_type, "s")
        self.assertNotIn("\x07", ws.cell(row=2, column=HEADERS.index("B") + 1).value)

    def test_a_select_2_question_and_an_empty_e_column(self):
        many = next(q for q in self.questions if q["select"] == 2)
        one = next(q for q in self.questions if q["select"] == 1)
        rows = list(load_workbook(self.m.make(1, [many, one])).active.iter_rows(values_only=True))
        a, b = dict(zip(HEADERS, rows[1])), dict(zip(HEADERS, rows[2]))
        self.assertEqual(a["Correct answer"], "A, B")
        self.assertTrue(a["E"])
        self.assertIn(b["E"], (None, ""))

    def test_a_missing_file_and_changed_columns_are_errors(self):
        with self.assertRaises(MarkingError):
            self.m.read(2)
        path = self.m.make(1, self.questions[:2])
        wb = load_workbook(path)
        wb.active.delete_cols(3)
        wb.save(path)
        with self.assertRaises(MarkingError) as e:
            self.m.read(1)
        self.assertIn("columns", str(e.exception))

    def test_the_backslash_of_the_records_is_not_shown(self):
        self.assertEqual(plain("e.g., \\*\\*/\\*.test.tsx"), "e.g., **/*.test.tsx")
        self.assertEqual(plain("--resume \\<session-name\\>"), "--resume <session-name>")

    def test_bad_marks_for_each_scenario_and_each_domain(self):
        s1, s2 = PLAN["slots"][0], PLAN["slots"][20]
        marks = [{"question_id": s1["slot_id"] + "-T1", "mark": "bad", "reason": "x"},
                 {"question_id": s2["slot_id"] + "-T2", "mark": "bad", "reason": "y"},
                 {"question_id": PLAN["slots"][5]["slot_id"] + "-T1", "mark": "good", "reason": ""}]
        out = self.m.bad_marks_by_segment(marks)
        domain = {t["id"]: t["domain"] for t in RECORDS["task_statements"]}
        expected_s, expected_d = {}, {}
        for s in (s1, s2):
            expected_s[s["scenario"]] = expected_s.get(s["scenario"], 0) + 1
            d = domain[s["task_statement"]]
            expected_d[d] = expected_d.get(d, 0) + 1
        self.assertEqual(out["scenario"], dict(sorted(expected_s.items())))
        self.assertEqual(out["domain"], dict(sorted(expected_d.items())))


if __name__ == "__main__":
    unittest.main()
