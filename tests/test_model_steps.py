"""Tests of the model steps 2, 4, 5 and 6, of the adapter and of the whole
pipeline. A fake client takes the place of the API. No test makes a model call.

Run from the folder ccar-f-mock-generator:
    python3 -m unittest discover tests
"""
import json
import os
import re
import shutil
import sys
import tempfile
import unittest
from types import SimpleNamespace

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))

from control import CallFailed, Control, RunStopped  # noqa: E402
from model_call import ModelCall, strict  # noqa: E402
from step1_plan import make_plan  # noqa: E402
from step2_write import Writer, request_text, examples, check_answer, give_letters  # noqa: E402
from step4_check_model import ModelCheck, FeatureSearch, show, guide_text  # noqa: E402
from step6_check_scenario import ScenarioCheck, code_count  # noqa: E402
from step7_marking import HEADERS  # noqa: E402
from steps import RealSteps, guide_points, load_records  # noqa: E402

RECORDS = load_records()
PLAN = make_plan(RECORDS, 1, 11, "2026-10-07")


def response(answer, stop="end_turn", model="claude-opus-5-5"):
    text = answer if isinstance(answer, str) else json.dumps(answer)
    return SimpleNamespace(content=[SimpleNamespace(type="text", text=text)], stop_reason=stop, model=model,
                           stop_details=SimpleNamespace(category="cyber"))


class FakeClient:
    """Answers each request with the next item of a list, or with a function."""

    def __init__(self, answer):
        self.answer = answer
        self.requests = []
        self.messages = self

    def stream(self, **kwargs):
        self.requests.append(kwargs)
        out = self.answer(kwargs) if callable(self.answer) else self.answer.pop(0)
        if isinstance(out, Exception):
            raise out
        return FakeStream(out)


class FakeStream:
    def __init__(self, message):
        self.message = message

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self.message


# Classes with the names of the SDK errors (the SDK is not installed for the tests).
class RateLimitError(Exception):
    pass


class BadRequestError(Exception):
    status_code = 400


class AuthenticationError(Exception):
    status_code = 401


class APIStatusError(Exception):
    status_code = 529


def written(slot, n_correct=None, tag=""):
    """A question that passes the rules of step 3. A correct option has the word RIGHT."""
    n_correct = slot["select"] if n_correct is None else n_correct
    options = []
    for i in range(slot["options"]):
        right = i < n_correct
        options.append({"text": "%s option number %d of the slot %s%s is here." % ("RIGHT" if right else "Another", i, slot["slot_id"], tag),
                        "correct": right, "reason": "The reason of option %d." % i})
    return {"status": "written",
            "question": ("A team builds an agent for the slot %s%s and sees a problem in production. The coordinator gets an empty "
                         "result from one subagent and writes an incomplete report for the user. Which change to the design "
                         "solves this problem best?") % (slot["slot_id"], tag),
            "options": options, "cannot_write_reason": None}


SLOT = PLAN["slots"][0]


class TestModelCall(unittest.TestCase):
    def test_the_form_is_sent_as_structured_outputs_with_no_tool(self):
        client = FakeClient([response({"x": 1})])
        answer, model = ModelCall(client).ask("system", "text", {"type": "object", "properties": {}})
        req = client.requests[0]
        self.assertEqual(req["output_config"]["format"]["type"], "json_schema")
        self.assertNotIn("tools", req)
        self.assertNotIn("tool_choice", req)
        self.assertEqual((answer, model), ({"x": 1}, "claude-opus-5-5"))

    def test_strict_adds_additional_properties_to_each_object(self):
        s = strict({"type": "object", "properties": {"a": {"type": "array", "items": {"type": "object", "properties": {}}}}})
        self.assertFalse(s["additionalProperties"])
        self.assertFalse(s["properties"]["a"]["items"]["additionalProperties"])

    def test_each_failure_that_a_new_call_can_repair_is_a_call_that_failed(self):
        for out in (response({}, stop="refusal"), response({}, stop="max_tokens"), response("not json"),
                    RateLimitError("429"), APIStatusError("overloaded"),
                    SimpleNamespace(content=[], stop_reason="end_turn", model="m")):
            with self.assertRaises(CallFailed):
                ModelCall(FakeClient([out])).ask("s", "t", {})

    def test_an_error_that_a_new_call_cannot_repair_stops_the_run(self):
        for out in (BadRequestError("schema"), AuthenticationError("no key"), TypeError("wrong parameter")):
            with self.assertRaises(RunStopped):
                ModelCall(FakeClient([out])).ask("s", "t", {})

    def test_the_call_streams_with_a_high_limit(self):
        client = FakeClient([response({"x": 1})])
        ModelCall(client).ask("s", "t", {})
        self.assertGreaterEqual(client.requests[0]["max_tokens"], 32000)

    def test_a_thinking_block_before_the_text(self):
        msg = response({"status": "Pass"})
        msg.content.insert(0, SimpleNamespace(type="thinking", thinking=""))
        answer, _ = ModelCall(FakeClient([msg])).ask("s", "t", {})
        self.assertEqual(answer, {"status": "pass"})

    def test_enum_values_get_one_fixed_form(self):
        answer, _ = ModelCall(FakeClient([response({"results": [{"rule": "b1", "result": "Cannot_Decide", "letter": "c"}],
                                                     "selected": ["a", "b"], "status": "Answered"})])).ask("s", "t", {})
        self.assertEqual(answer, {"results": [{"rule": "B1", "result": "cannot_decide", "letter": "C"}],
                                  "selected": ["A", "B"], "status": "answered"})

    def test_the_reason_fields_ask_for_a_short_explanation(self):
        from step2_write import SCHEMA
        from step4_check_model import SCHEMA_4A
        self.assertIn("short explanation", SCHEMA["properties"]["options"]["items"]["properties"]["reason"]["description"])
        self.assertIn("short explanation", SCHEMA_4A["properties"]["reason"]["description"])


class TestWriter(unittest.TestCase):
    def request(self, step=2, **extra):
        r = {"slot": SLOT, "point": SLOT["point"], "step": step, "shown_try": None, "failed_question": None,
             "findings": [], "other_questions": []}
        r.update(extra)
        return r

    def test_the_request_has_each_part_of_decision_0029(self):
        text = request_text(self.request(), RECORDS)
        self.assertIn("A11. The question is not a negative question.", text)
        self.assertIn("B6.", text)
        self.assertIn("B7. Each actor has one name", text)
        self.assertIn("Use one name for one actor.", text)
        self.assertIn(">>> %s:" % SLOT["point"], text)
        scen = next(s for s in RECORDS["scenarios"] if s["id"] == SLOT["scenario"])
        self.assertIn(scen["story"][:80], text)
        for t in RECORDS["section_17"]["technologies"]:
            self.assertIn(re.sub(r"\\(.)", r"\1", t)[:40], text)
        self.assertNotIn("A NEW TRY", text)

    def test_a_two_answer_slot_asks_for_two_different_parts(self):
        # Decision 0054.
        two = next(x for x in PLAN["slots"] if x["select"] == 2)
        one = next(x for x in PLAN["slots"] if x["select"] == 1)
        sentence = "The two correct options test two different parts of the point."
        self.assertIn(sentence, request_text(self.request(slot=two, point=two["point"]), RECORDS))
        self.assertNotIn(sentence, request_text(self.request(slot=one, point=one["point"]), RECORDS))
        self.assertIn("A12. When the slot has", request_text(self.request(), RECORDS))

    def test_three_examples_from_other_scenarios_with_no_letters(self):
        ex = examples(RECORDS, SLOT)
        self.assertEqual(len(ex), 3)
        samples = {s["question"]: s for s in RECORDS["sample_questions"]}
        for e in ex:
            self.assertNotEqual(samples[e["question"]]["scenario"], SLOT["scenario"])
            self.assertEqual(sum(o["correct"] for o in e["options"]), 1)
            self.assertTrue(all(set(o) == {"text", "correct"} for o in e["options"]))
        self.assertEqual(examples(RECORDS, SLOT), ex)

    def test_step_5_adds_the_failed_question_its_findings_and_the_other_question(self):
        failed = dict(written(SLOT), question_id="M1-S01-T1", select=SLOT["select"])
        failed["options"] = give_letters(failed, SLOT)
        other = {"question_id": "M1-S02-T1", "question": "Another question?", "options": [{"text": "x"}]}
        f = {"step": 6, "rule": "D1", "detail": "Same scene.", "letter": None, "point": None, "other_questions": ["M1-S02-T1"]}
        text = request_text(self.request(5, failed_question=failed, findings=[f], other_questions=[other]), RECORDS)
        self.assertIn("A NEW TRY", text)
        self.assertIn('"letter": "A"', text)
        self.assertIn("Same scene.", text)
        self.assertIn("Another question?", text)

    def test_the_answer_gets_letters_from_code(self):
        client = FakeClient([response(written(SLOT))])
        out = Writer(RECORDS, ModelCall(client)).write(self.request())
        self.assertEqual([o["letter"] for o in out["options"]], list("ABCDE"[:SLOT["options"]]))
        self.assertEqual(out["model"], "claude-opus-5-5")
        again = Writer(RECORDS, ModelCall(FakeClient([response(written(SLOT))]))).write(self.request())
        self.assertEqual(out["options"], again["options"])

    def test_cannot_write(self):
        answer = {"status": "cannot_write", "question": None, "options": [], "cannot_write_reason": "It does not fit."}
        out = Writer(RECORDS, ModelCall(FakeClient([response(answer)]))).write(self.request())
        self.assertEqual((out["status"], out["cannot_write_reason"]), ("cannot_write", "It does not fit."))

    def test_an_answer_that_is_not_complete_is_a_call_that_failed(self):
        for bad in ({"status": "maybe"}, {"status": "cannot_write", "cannot_write_reason": ""},
                    {"status": "written", "question": "", "options": []}):
            with self.assertRaises(CallFailed):
                check_answer(bad)


def question_of(slot, answer=None):
    a = answer or written(slot)
    return {"question_id": slot["slot_id"] + "-T1", "slot_id": slot["slot_id"], "try": 1, "point": slot["point"],
            "select": slot["select"], "status": "written", "question": a["question"], "options": give_letters(a, slot)}


def blind_for(question, status="answered"):
    letters = [o["letter"] for o in question["options"] if "RIGHT" in o["text"]]
    return {"status": status, "selected": letters if status == "answered" else [], "reason": "My reason."}


def check_4b(results=None, features=()):
    rs = results or {}
    return {"results": [{"rule": r, "result": rs.get(r, ("pass",))[0], "finding": rs.get(r, (None, None))[1],
                         "letter": None, "point": None} for r in ["B1", "B2", "B3", "B4", "B5", "B6", "B7", "B8", "B9"]],
            "features": list(features)}


class TestModelCheck(unittest.TestCase):
    def run_check(self, blind, check, q=None):
        q = q or question_of(SLOT)
        client = FakeClient([response(blind), response(check)])
        return ModelCheck(RECORDS, ModelCall(client)).check(q, SLOT), client

    def test_the_blind_call_sees_no_marks_and_no_reasons(self):
        q = question_of(SLOT)
        v, client = self.run_check(blind_for(q), check_4b(), q)
        blind_text = client.requests[0]["messages"][0]["content"]
        self.assertIn("Select %d." % SLOT["select"], blind_text)
        self.assertNotIn("Marked", blind_text)
        self.assertNotIn("The reason of option", blind_text)
        self.assertIn("Marked: correct", client.requests[1]["messages"][0]["content"])
        self.assertEqual(v["verdict"], "pass")

    def test_the_blind_call_gets_the_guide_text_and_the_story(self):
        q = question_of(SLOT)
        _, client = self.run_check(blind_for(q), check_4b(), q)
        blind_text = client.requests[0]["messages"][0]["content"]
        scen = next(s for s in RECORDS["scenarios"] if s["id"] == SLOT["scenario"])
        self.assertIn(scen["story"][:80], blind_text)
        self.assertEqual(len(re.findall(r"^  \d\.\d-[KS]\d+:", blind_text, re.M)), 240)

    def test_the_rule_check_gets_the_rules_the_slot_and_the_point(self):
        q = question_of(SLOT)
        _, client = self.run_check(blind_for(q), check_4b(), q)
        text = client.requests[1]["messages"][0]["content"]
        for part in ("B1.", "B6.", "B7.", "B8.", "point %s" % SLOT["point"], "task statement %s" % SLOT["task_statement"]):
            self.assertIn(part, text)
        # Decision 0055: the 12 sample questions, each with its correct answer.
        self.assertEqual(len(re.findall(r"^Sample question \d+ \(scenario \d\):", text, re.M)), 12)
        self.assertEqual(text.count("\nCorrect: "), 12)

    def test_the_prompt_names_the_right_number_of_rules(self):
        import step4_check_model as s4
        words = {6: "six", 7: "seven", 8: "eight", 9: "nine"}
        self.assertIn("against %s rules" % words[len(s4.RULES_B)], s4.SYSTEM_4B)

    def test_the_form_of_call_4b_has_each_rule(self):
        # The review of 8 October 2026: the enum of the form must have each rule,
        # or each real call 4b fails.
        import step4_check_model as s4
        rule = s4.SCHEMA_4B["properties"]["results"]["items"]["properties"]["rule"]
        self.assertEqual(rule["enum"], s4.RULES_B)
        self.assertIn("B8.", s4.RULES_TEXT)

    def test_b6_finds_a_flag_with_its_hyphens(self):
        # Decision 0061: the guide names -p and --print. A one-letter flag that
        # the guide does not name stays not named.
        from step4_check_model import FeatureSearch
        f = FeatureSearch(RECORDS)
        for name in ("-p", "--print", "`-p`", "--json-schema", "--output-format", "--resume"):
            self.assertTrue(f.named(name), name)
        for name in ("-x", "-r", "p", "--printx"):
            self.assertFalse(f.named(name), name)

    def test_the_rule_check_gets_the_complete_guide_text(self):
        text = guide_text(RECORDS, SLOT["scenario"])
        self.assertEqual(len(re.findall(r"^  \d\.\d-[KS]\d+:", text, re.M)), 240)

    def test_a_different_blind_answer_fails_b1(self):
        q = question_of(SLOT)
        wrong = {"status": "answered", "reason": "x",
                 "selected": [o["letter"] for o in q["options"] if "RIGHT" not in o["text"]][:SLOT["select"]]}
        v, _ = self.run_check(wrong, check_4b(), q)
        self.assertEqual(v["verdict"], "fail")
        self.assertEqual(v["findings"][0]["rule"], "B1")

    def test_cannot_decide(self):
        q = question_of(SLOT)
        self.assertEqual(self.run_check(blind_for(q, "cannot_decide"), check_4b(), q)[0]["verdict"], "cannot_decide")
        v, _ = self.run_check(blind_for(q), check_4b({"B4": ("cannot_decide", "Unclear point.")}), q)
        self.assertEqual(v["verdict"], "cannot_decide")
        v, _ = self.run_check(blind_for(q), check_4b({"B4": ("cannot_decide", "x"), "B2": ("fail", "y")}), q)
        self.assertEqual(v["verdict"], "fail")

    def test_b6_code_searches_the_feature_names(self):
        q = question_of(SLOT)
        v, _ = self.run_check(blind_for(q), check_4b(features=["tool-choice", "CLAUDE.md", "fork session"]), q)
        self.assertEqual(v["verdict"], "pass", v["findings"])
        v, _ = self.run_check(blind_for(q), check_4b(features=["--turbo-mode"]), q)
        self.assertEqual(v["verdict"], "fail")
        self.assertIn("--turbo-mode", v["findings"][0]["detail"])
        self.assertFalse([f for f in v["features"] if f["found"]])

    def test_answers_that_are_not_valid_are_calls_that_failed(self):
        q = question_of(SLOT)
        too_many = {"status": "answered", "selected": ["A", "B", "C"], "reason": "x"}
        bad_letter = check_4b()
        bad_letter["results"][0].update(result="fail", finding="x", letter="E" if SLOT["options"] == 4 else "Z")
        bad_point = check_4b()
        bad_point["results"][1].update(result="fail", finding="x", point="9.9-K9")
        missing = check_4b()
        missing["results"].pop()
        no_finding = check_4b({"B3": ("fail", "")})
        twice = {"status": "answered", "selected": ["A"] * SLOT["select"], "reason": "x"}
        unknown = {"status": "answered", "selected": ["E" if SLOT["options"] == 4 else "Z"] + ["A"] * (SLOT["select"] - 1), "reason": "x"}
        cases = [(too_many, check_4b()), (blind_for(q), bad_letter), (blind_for(q), bad_point),
                 (blind_for(q), missing), (blind_for(q), no_finding)]
        if SLOT["select"] == 2:
            cases.append((twice, check_4b()))
        cases.append((unknown, check_4b()))
        for blind, check in cases:
            with self.assertRaises(CallFailed):
                self.run_check(blind, check, q)

    def test_a_cannot_decide_answer_has_no_letters_and_is_valid(self):
        q = question_of(SLOT)
        self.assertEqual(self.run_check({"status": "cannot_decide", "selected": [], "reason": "x"}, check_4b(), q)[0]["verdict"],
                         "cannot_decide")

    def test_feature_search_rules(self):
        fs = FeatureSearch(RECORDS)
        self.assertTrue(fs.named("TOOL_CHOICE"))
        self.assertTrue(fs.named("`tool_choice`"))
        self.assertTrue(fs.named("**/*.test.tsx"))
        self.assertFalse(fs.named("teleport_agent"))
        self.assertFalse(fs.named("  "))

    def test_a_name_inside_a_different_word_is_not_named(self):
        fs = FeatureSearch(RECORDS)
        self.assertFalse(fs.named("rovision"))     # only a part of the word "provisions"
        # "Vision" is a complete word on the list "out of scope" of Section 17. If such a
        # name counts as named is an open item (decision 0030). Now it counts.
        self.assertTrue(fs.named("Vision"))
        self.assertFalse(fs.named("-c"))
        self.assertFalse(fs.named("e"))
        self.assertTrue(fs.named("isError"))


class TestScenarioCheck(unittest.TestCase):
    def questions(self, scenario):
        return [question_of(s) for s in PLAN["slots"] if s["scenario"] == scenario]

    def test_two_calls_and_all_rules(self):
        def answer(req):
            text = req["messages"][0]["content"]
            rules = ["D1", "D4"] if "D1." in text else ["D2", "D3"]
            return response({"results": [{"rule": r, "result": "pass", "findings": []} for r in rules]})
        client = FakeClient(answer)
        sc = PLAN["scenarios"][0]
        out = ScenarioCheck(PLAN, ModelCall(client)).check(sc, self.questions(sc))
        self.assertEqual(sorted(r["rule"] for r in out), ["D1", "D2", "D3", "D4"])
        self.assertEqual(len(client.requests), 2)
        self.assertNotIn("RIGHT option", client.requests[1]["messages"][0]["content"])

    def test_c3_and_c2(self):
        sc = PLAN["scenarios"][0]
        qs = self.questions(sc)
        qs[1]["question"] = qs[0]["question"]
        self.assertEqual(code_count(sc, qs, PLAN["slots"])[0]["rule"], "C3")
        with self.assertRaises(RunStopped):
            code_count(sc, qs[:-1], PLAN["slots"])

    def test_c3_compares_the_other_scenarios(self):
        # Decision 0052: the same text in two scenarios is a finding of C3.
        sc, other = PLAN["scenarios"][0], PLAN["scenarios"][1]
        qs, os_ = self.questions(sc), self.questions(other)
        self.assertEqual(code_count(sc, qs, PLAN["slots"], os_), [])
        os_[2]["question"] = "  " + qs[4]["question"].upper()
        out = code_count(sc, qs, PLAN["slots"], os_)
        self.assertEqual(out[0]["findings"][0]["question_ids"], [qs[4]["question_id"], os_[2]["question_id"]])
        self.assertIn("scenario %s" % other, out[0]["findings"][0]["detail"])

    def test_each_result_of_a_call_keeps_the_model(self):
        def answer(req):
            rules = ["D1", "D4"] if "D1." in req["messages"][0]["content"] else ["D2", "D3"]
            return response({"results": [{"rule": r, "result": "pass", "findings": []} for r in rules]}, model="claude-x")
        sc = PLAN["scenarios"][0]
        out = ScenarioCheck(PLAN, ModelCall(FakeClient(answer))).check(sc, self.questions(sc))
        self.assertTrue(all(r["model"] == "claude-x" for r in out if r["rule"].startswith("D")))


class TestWholePipeline(unittest.TestCase):
    """All real steps, from the plan to the two files of the mock. The fake
    client answers as a good writer and good checkers do."""

    def setUp(self):
        self.root = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.root)

    def test_from_the_plan_to_the_files(self):
        def answer(req):
            system, text = req["system"], req["messages"][0]["content"]
            if system.startswith("You write"):
                slot_id = re.search(r'"point": "([^"]+)"', text.split("The slot: ")[1]).group(1)
                slot = next(s for s in plan_slots if s["point"] == slot_id)
                return response(written(slot))
            if system.startswith("You answer"):
                select = int(re.search(r"Select (\d)\.", text).group(1))
                letters = re.findall(r"^([A-E])\. RIGHT", text, re.M)
                assert len(letters) == select
                return response({"status": "answered", "selected": letters, "reason": "x"})
            if system.startswith("You check one exam question"):
                return response(check_4b())
            rules = ["D1", "D4"] if "D1." in text else ["D2", "D3"]
            return response({"results": [{"rule": r, "result": "pass", "findings": []} for r in rules]})
        run = os.path.join(self.root, "runs", "mock-01")
        steps = RealSteps(run, RECORDS, ModelCall(FakeClient(answer)), os.path.join(self.root, "marking"),
                          os.path.join(self.root, "mocks"), today=lambda: "2026-10-07")
        c = Control(run, 1, 11, steps, guide_points(RECORDS), today=lambda: "2026-10-07")
        plan_slots = steps.plan(1, 11)["slots"]
        report = c.run()
        self.assertEqual(report["state"], "ready_to_mark", report["waiting_for"])
        path = os.path.join(self.root, "marking", "M1-round1.xlsx")
        wb = load_workbook(path)
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            row[HEADERS.index("Mark")].value = "good"
        wb.save(path)
        self.assertEqual(c.read_marks()["waiting_for"], ["accept the mock"])
        self.assertEqual(c.accept()["state"], "built")
        with open(os.path.join(self.root, "mocks", "mock-01.questions.json"), encoding="utf-8") as f:
            built = json.load(f)
        self.assertEqual(sum(len(g["questions"]) for g in built["groups"]), 60)


if __name__ == "__main__":
    unittest.main()
