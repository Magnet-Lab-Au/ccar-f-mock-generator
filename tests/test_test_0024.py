"""Tests of the script of the test before the first full run (decision 0042).
A fake client takes the place of the API. No test makes a model call."""
import json
import os
import shutil
import sys
import tempfile
import unittest
from collections import Counter
from unittest import mock

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import test_0024 as t  # noqa: E402
from model_call import ModelCall  # noqa: E402
from step1_plan import make_plan  # noqa: E402
from step7_marking import HEADERS  # noqa: E402
from steps import load_records  # noqa: E402
from test_model_steps import FakeClient, response, written, check_4b  # noqa: E402

RECORDS = load_records()
PLAN = make_plan(RECORDS, 1, 366341, "2026-10-07")
DOMAIN = {x["id"]: x["domain"] for x in RECORDS["task_statements"]}


class TestSelect(unittest.TestCase):
    def test_ten_slots_cover_scenarios_domains_and_two_multiple_response(self):
        slots = t.select_slots(PLAN, RECORDS)
        self.assertEqual(len(slots), 10)
        self.assertEqual({s["scenario"] for s in slots}, set(PLAN["scenarios"]))
        self.assertEqual({DOMAIN[s["task_statement"]] for s in slots}, {1, 2, 3, 4, 5})
        self.assertEqual(Counter(s["type"] for s in slots)["multiple_response"], 2)
        self.assertEqual(t.select_slots(PLAN, RECORDS), slots)


class TestParts(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp()
        self.patches = [mock.patch.object(t, "FOLDER", os.path.join(self.root, "runs")),
                        mock.patch.object(t, "MARKING", os.path.join(self.root, "marking")),
                        mock.patch.object(t, "load_plan", lambda: PLAN)]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()
        shutil.rmtree(self.root)

    def test_part1_writes_checks_and_makes_a_marking_file(self):
        def answer(req):
            point = json.loads(req["messages"][0]["content"].split("The slot: ")[1].split("\n")[0])["point"]
            slot = next(s for s in PLAN["slots"] if s["point"] == point)
            return response(written(slot))
        mc = ModelCall(FakeClient(answer))
        r = t.part1(RECORDS, PLAN, mc)
        self.assertEqual(len(r["questions"]), 10)
        self.assertTrue(all(e["step3"]["verdict"] == "pass" for e in r["questions"]))
        self.assertTrue(r["marking_file"].endswith("test-0024-round1.xlsx"))
        self.assertEqual(r["cost"]["calls"], 0)   # the fake response has no usage

    def test_part2_counts_the_right_answers(self):
        def answer(req):
            text = req["messages"][0]["content"]
            sample = next(s for s in RECORDS["sample_questions"] if s["question"][:60] in text)
            return response({"status": "answered", "selected": sample["correct_answer"], "reason": "x"})
        r = t.part2(RECORDS, ModelCall(FakeClient(answer)))
        self.assertEqual((r["right"], r["of"]), (12, 12))

    def test_part3_good_must_pass_and_damaged_must_fail_on_its_rule(self):
        def write_answer(req):
            point = json.loads(req["messages"][0]["content"].split("The slot: ")[1].split("\n")[0])["point"]
            slot = next(s for s in PLAN["slots"] if s["point"] == point)
            return response(written(slot))
        t.part1(RECORDS, PLAN, ModelCall(FakeClient(write_answer)))
        path = os.path.join(self.root, "marking", "test-0024-round1.xlsx")
        wb = load_workbook(path)
        for row in wb.active.iter_rows(min_row=2):
            row[HEADERS.index("Mark")].value = "good"
        wb.save(path)
        ids = [e["question"]["question_id"] for e in t.load("part1.json")["questions"]]
        t.save("damaged.json", [{"question_id": ids[0], "rule": "B1", "what": "code moves the mark"},
                                {"question_id": ids[1], "rule": "B5", "what": "a scene outside the story",
                                 "question": {"question": "A scene on the moon?"}}])

        def check_answer(req):
            text = req["messages"][0]["content"]
            if "A scene on the moon?" in text:
                return response(check_4b({"B5": ("fail", "Outside the story.")}))
            lines = text.splitlines()
            moved = any("RIGHT option" in lines[i] and "Marked: wrong" in lines[i + 1] for i in range(len(lines) - 1))
            return response(check_4b({"B1": ("fail", "The mark is wrong.")}) if moved else check_4b())
        r = t.part3(RECORDS, ModelCall(FakeClient(check_answer)))
        self.assertEqual(r["of"], 12)
        self.assertEqual(r["as_expected"], 12, [e for e in r["checks"] if not e["as_expected"]])


if __name__ == "__main__":
    unittest.main()
