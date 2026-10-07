"""Step 7 of the pipeline: the marking file of the owner (decisions 0015, 0022, 0035 and 0039).

Code builds one Excel file for each round of marking, for example
marking/M1-round1.xlsx. It has one row for each question. It shows no verdict
of a checker. The owner fills two columns: "Mark" (good or bad) and "Reason".

The control decides which questions go into a round, checks that each mark is
complete, and saves the marks record (decision 0039). This step only writes
and reads the file.
"""
import os
import re
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["Question ID", "Scenario", "Point", "Point text", "Question", "Select",
           "A", "B", "C", "D", "E", "Correct answer", "Reasons of the options",
           "Mark", "Reason (needed for a bad mark)"]
SHEET = "Marks"
FILL = PatternFill("solid", fgColor="FFF2CC")   # the yellow of the columns that the owner fills
WIDTHS = [12, 22, 9, 40, 60, 9, 30, 30, 30, 30, 30, 10, 60, 10, 40]


class MarkingError(Exception):
    """The step cannot write or read the file. The message names the cause."""


def plain(text):
    """The records file keeps a backslash from the format of the guide, for
    example "\\*\\*/\\*.test.tsx". The owner reads the text without it."""
    return re.sub(r"\\(.)", r"\1", text)


def _cell_text(value):
    """Text for a cell. Remove characters that Excel does not permit."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


class Marking:
    def __init__(self, plan, records, folder):
        self.plan = plan
        self.prefix = "M%d" % plan["mock"]
        self.folder = folder
        self.slots = {s["slot_id"]: s for s in plan["slots"]}
        self.scenarios = {s["id"]: s["name"] for s in records["scenarios"]}
        self.points = {p["id"]: p["text"] for t in records["task_statements"] for p in t["points"]}
        self.domain_of = {t["id"]: t["domain"] for t in records["task_statements"]}

    def path(self, round_number):
        return os.path.join(self.folder, "%s-round%d.xlsx" % (self.prefix, round_number))

    def make(self, round_number, questions):
        """Write the marking file of one round. Never write over a file. If the
        same file exists already with no marks, for example after a stop of the
        program, return it."""
        path = self.path(round_number)
        if os.path.exists(path):
            ids = [q["question_id"] for q in questions]
            try:
                marks = self.read(round_number)
            except MarkingError:
                marks = None
            if marks is not None and [m["question_id"] for m in marks] == ids \
                    and not any(m["mark"] or m["reason"] for m in marks):
                return path
            raise MarkingError("The marking file exists already: %s" % path)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(HEADERS)
        for q in questions:
            slot = self.slots[q["slot_id"]]
            point = q.get("point", slot["point"])
            letters = {o["letter"]: o["text"] for o in q["options"]}
            ws.append([_cell_text(v) for v in [q["question_id"],
                       "%d. %s" % (slot["scenario"], self.scenarios[slot["scenario"]]),
                       point, plain(self.points.get(point, "")),
                       q["question"], "Select %d." % q["select"]]
                      + [letters.get(x, "") for x in "ABCDE"]
                      + [", ".join(o["letter"] for o in q["options"] if o["correct"]),
                         "\n".join("%s: %s" % (o["letter"], o["reason"]) for o in q["options"]),
                         None, None]])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"   # a text, not a formula
        mark_col, reason_col = HEADERS.index("Mark") + 1, HEADERS.index("Reason (needed for a bad mark)") + 1
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            row[mark_col - 1].fill = FILL
            row[reason_col - 1].fill = FILL
            if row[0].row > 1:
                row[mark_col - 1].protection = Protection(locked=False)
                row[reason_col - 1].protection = Protection(locked=False)
        # The sheet is protected with no password. Only the two columns of the
        # owner can change, and sorting is off, so that a mark cannot move to a
        # different question. The owner can remove the protection on purpose.
        ws.protection.sheet = True
        ws.protection.sort = True
        ws.protection.autoFilter = True
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "B2"
        if ws.max_row > 1:
            letter = ws.cell(row=1, column=mark_col).column_letter
            check = DataValidation(type="list", formula1='"good,bad"', allow_blank=True, showErrorMessage=True,
                                   errorTitle="Mark", error='Type "good" or "bad".')
            check.add("%s2:%s%d" % (letter, letter, ws.max_row))
            ws.add_data_validation(check)
        os.makedirs(self.folder, exist_ok=True)
        wb.save(path)
        return path

    def read(self, round_number):
        """Return the marks of the file: [{question_id, mark, reason}]. A mark that
        is missing or not "good" or "bad" comes back as it is, so that the control
        can name the row."""
        path = self.path(round_number)
        if not os.path.exists(path):
            raise MarkingError("No marking file: %s" % path)
        try:
            wb = load_workbook(path, read_only=True)
        except Exception as e:   # a damaged file
            raise MarkingError("The marking file cannot be read: %s" % e)
        if SHEET not in wb.sheetnames:
            raise MarkingError('The marking file has no sheet "%s".' % SHEET)
        rows = list(wb[SHEET].iter_rows(values_only=True))
        if not rows or list(rows[0][:len(HEADERS)]) != HEADERS:
            raise MarkingError("The columns of the marking file changed. Use the file as code made it.")
        i_id, i_mark, i_reason = (HEADERS.index("Question ID"), HEADERS.index("Mark"),
                                  HEADERS.index("Reason (needed for a bad mark)"))
        marks, seen = [], set()
        for row in rows[1:]:
            if not row or row[i_id] is None:
                continue
            qid = str(row[i_id]).strip()
            if qid in seen:
                raise MarkingError("The question %s is in the marking file two times." % qid)
            seen.add(qid)
            for value in (row[i_mark], row[i_reason]):
                if isinstance(value, str) and value.startswith("="):
                    raise MarkingError("The row of %s has a formula. Type the mark and the reason as text." % qid)
            mark = row[i_mark]
            mark = mark.strip().lower() if isinstance(mark, str) and mark.strip() else None
            reason = row[i_reason]
            reason = str(reason).strip() if reason is not None and str(reason).strip() else ""
            marks.append({"question_id": qid, "mark": mark, "reason": reason})
        return marks

    def bad_marks_by_segment(self, marks):
        """The bad marks for each scenario and each domain (decision 0035, part 4)."""
        scenario, domain = Counter(), Counter()
        for m in marks:
            if m["mark"] != "bad":
                continue
            slot = self.slots[m["question_id"].rsplit("-T", 1)[0]]
            scenario[slot["scenario"]] += 1
            domain[self.domain_of[slot["task_statement"]]] += 1
        return {"scenario": dict(sorted(scenario.items())), "domain": dict(sorted(domain.items()))}
